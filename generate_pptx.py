#!/usr/bin/env python3
"""Generate presentation slides for Challenge 2."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
AMBER = RGBColor(0xD9, 0x77, 0x06)
DARK = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x6B, 0x72, 0x80)

def add_slide(prs, title_text, bullets, layout_idx=1):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    title = slide.shapes.title
    title.text = title_text
    for run in title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = DARK
        run.font.bold = True
        run.font.size = Pt(28)

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(16)
        p.font.color.rgb = DARK
        p.space_after = Pt(6)
    return slide

def add_title_slide(prs, title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title_text
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = DARK
        run.font.bold = True
        run.font.size = Pt(32)
    slide.placeholders[1].text = subtitle_text
    return slide

def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 — Title
    add_title_slide(prs, 'Challenge 2 — Home Handyman Portfolio',
                    'COMP5450 Mobile Programming · Group 5\nMay 27, 2026')

    # 2 — Problem & Goal
    add_slide(prs, 'Problem & Goal', [
        'Homeowners need a single, trustworthy view of a handyman\'s work.',
        'Our app combines services, portfolio, reviews, and contact in one place.',
        'The handout says: "showcase skills that attract potential customers."',
        'Cross-platform Flutter lets us ship to Android, iOS, and web from one codebase.',
    ])

    # 3 — Tech Stack
    add_slide(prs, 'Tech Stack', [
        'Flutter 3.32 + Dart — cross-platform framework from the course',
        'Material 3 design with a warm amber/orange theme',
        'Repository pattern for data access (program to an interface)',
        'Firebase-ready: abstract PortfolioRepository → swap mock for Firestore',
        'Tested with flutter analyze (0 issues) and 7 widget tests',
    ])

    # 4 — Architecture
    add_slide(prs, 'Architecture', [
        'models/ — Service, Project, Review (plain Dart classes)',
        'data/ — PortfolioRepository (abstract) + MockPortfolioRepository',
        'screens/ — one file per screen (Home, Services, Portfolio, About, Reviews, Contact)',
        'widgets/ — reusable cards (HeroBanner, ServiceCard, ProjectCard, ReviewCard, StatChip)',
        'Swapping to Firestore requires changing only one class — no screen edits.',
    ])

    # 5 — Home Screen
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Home Screen'
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = DARK
        run.font.bold = True
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for b in ['Hero banner with Call / Book / Portfolio buttons',
              'Stat chips: 12 years, 340+ projects, 85% repeat clients',
              'Featured services grid (top 4 of 6)',
              'ConstrainedBox(maxWidth: 600) for responsive centering']:
        p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(16)
    img_path = os.path.join(HERE, 'screenshot-home.png')
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(8.5), Inches(1.2), height=Inches(5.5))

    # 6 — Services & Portfolio
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Services & Portfolio'
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = DARK
        run.font.bold = True
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    for b in ['Services: 6 categories in a 2-column grid with icons',
              'Portfolio: project cards with category tags and before/after labels',
              'Placeholder images with color overlays (no network images)',
              'All data from MockPortfolioRepository → easy to swap']:
        p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
        p.text = b
        p.font.size = Pt(16)
    img_path = os.path.join(HERE, 'screenshot-services.png')
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(7.0), Inches(1.2), height=Inches(5.5))
    img_path2 = os.path.join(HERE, 'screenshot-portfolio.png')
    if os.path.exists(img_path2):
        slide.shapes.add_picture(img_path2, Inches(10.0), Inches(1.2), height=Inches(5.5))

    # 7 — Reviews & Contact
    add_slide(prs, 'Reviews & Contact', [
        'Reviews: average rating header + testimonial cards (name, stars, quote)',
        'Contact: business info tiles + message form with validation',
        'Form calls submitContact() — simulates 400ms delay, shows snackbar',
        'When Firebase is connected, the same method writes to Firestore',
    ])

    # 8 — About Screen
    add_slide(prs, 'About Screen', [
        'Accessible from the info icon in the app bar',
        'Profile avatar, name, tagline',
        'Stat chips: years, projects, repeat rate',
        'Bio paragraph, certifications, and service area',
    ])

    # 9 — Testing & Quality
    add_slide(prs, 'Testing & Quality', [
        'flutter analyze — 0 warnings or errors',
        'flutter test — 7 widget tests, all pass',
        'Tests cover: all 5 tabs, About page nav, form validation',
        'flutter build web — successful, used for screenshots',
    ])

    # 10 — Firebase Plan
    add_slide(prs, 'Firebase Integration Plan', [
        '1. Add firebase_core, cloud_firestore, firebase_storage to pubspec.yaml',
        '2. Create FirestorePortfolioRepository implementing the same interface',
        '3. Replace mock in main.dart — zero screen changes',
        '"Program to an interface, not an implementation"',
    ])

    # 11 — Conclusion
    add_slide(prs, 'Conclusion', [
        'Six-screen Flutter portfolio app for a home handyman service',
        'Material 3 with warm amber theme, responsive layout',
        'Repository pattern keeps data access swappable',
        'All tests pass, analysis clean, web build verified',
        'Firebase integration is a one-class swap',
        'Thank you — questions?',
    ])

    out_path = os.path.join(HERE, 'presentation',
                            'Challenge2_Group5_HandymanPortfolio.pptx')
    prs.save(out_path)
    print(f'✓ Generated {out_path}')

if __name__ == '__main__':
    build()
